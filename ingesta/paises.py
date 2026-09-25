"""Huéspedes por país de origen (solo Marriott Buenos Aires).

Se lee de la planilla "Venta x PAIS" que se descarga de Drive: la hoja con columnas
Continente / Año / PAÍS / Mes / N° Mes / Importe (el importe es la cantidad de huéspedes).

    python -m ingesta.run paises Venta_x_PAIS_BS_AS.xlsx
"""
from __future__ import annotations

import io
import unicodedata

from . import almacen

COLS = ["hotel", "mes", "pais", "continente", "huespedes"]
HOTEL = "marriott"


def _norm(v) -> str:
    t = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    return " ".join(t.upper().replace("°", "").split())


def _buscar_hoja(wb):
    for ws in wb.worksheets:
        for fila in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            cab = [_norm(c) for c in fila]
            if {"CONTINENTE", "ANO", "PAIS", "IMPORTE"} <= set(cab) and ("N MES" in cab or "MES" in cab):
                return ws, cab
    return None, None


def es_planilla(contenido: bytes) -> bool:
    from openpyxl import load_workbook

    ws, _ = _buscar_hoja(load_workbook(io.BytesIO(contenido), read_only=True, data_only=True))
    return ws is not None


def leer(contenido: bytes) -> list[dict]:
    from openpyxl import load_workbook

    ws, cab = _buscar_hoja(load_workbook(io.BytesIO(contenido), read_only=True, data_only=True))
    if ws is None:
        raise ValueError("No se encontró la hoja con Continente / Año / PAÍS / Mes / Importe")
    i = {k: cab.index(k) for k in ("CONTINENTE", "ANO", "PAIS", "IMPORTE")}
    i_mes = cab.index("N MES") if "N MES" in cab else cab.index("MES")
    total: dict[tuple, dict] = {}
    for fila in ws.iter_rows(values_only=True):
        anio, mes, n = fila[i["ANO"]], fila[i_mes], fila[i["IMPORTE"]]
        if not isinstance(anio, (int, float)) or not isinstance(mes, (int, float)) or not isinstance(n, (int, float)):
            continue
        if not n:
            continue
        pais = " ".join(str(fila[i["PAIS"]] or "").split()).upper()
        clave = (f"{int(anio)}-{int(mes):02d}", pais)
        x = total.setdefault(clave, {"hotel": HOTEL, "mes": clave[0], "pais": pais,
                                     "continente": " ".join(str(fila[i["CONTINENTE"]] or "").split()).upper(),
                                     "huespedes": 0})
        x["huespedes"] += n
    return sorted(total.values(), key=lambda r: (r["mes"], -r["huespedes"]))


def guardar(filas: list[dict]) -> int:
    """La planilla es acumulada: reemplaza todo lo anterior del hotel."""
    return almacen.reemplazar_bloque("paises", COLS, filas, ("hotel",))
