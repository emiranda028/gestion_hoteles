"""Persistencia en CSV (fuente de la web app) + Excel con el formato de la base de Power BI.

Tablas (todas en data/):
  hf.csv                una fila por hotel y día (History / Forecast), USD
  flash.csv             conceptos del Manager Flash por hotel y fecha de negocio (día / mes / año), USD
  pickup.csv            foto diaria del total del mes (on the books) para el mes en curso y los siguientes
  bonvoy.csv            llegadas de socios Bonvoy por nivel
  disponibilidades.csv  saldos resumidos por grupo y día (mismo esquema que "Disponibilidades PBI")
  bancos.csv            detalle por cuenta bancaria / inversión / caja
  tipo_cambio.csv       dólar BNA vendedor por día
  paises.csv            huéspedes por país y mes (Marriott BA)
"""
from __future__ import annotations

import csv
import json
from datetime import date, datetime, timezone
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
DATA = RAIZ / "data"
REGISTRO = DATA / "procesados.json"
EXCEL = DATA / "hoteles.xlsx"

HF_COLS = ["hotel", "fecha", "tipo", "total_occ", "arr_rooms", "comp_rooms", "house_use", "deduct_indiv",
           "deduct_group", "occ_pct", "room_revenue", "adr", "dep_rooms", "day_use", "no_show", "ooo", "personas",
           "fecha_reporte"]
FLASH_COLS = ["hotel", "fecha", "fecha_reporte", "concepto", "dia", "mes", "anio"]
PICKUP_COLS = ["hotel", "fecha_reporte", "mes", "noches", "grupo", "occ_pct", "revenue", "adr"]
BONVOY_COLS = ["hotel", "fecha", "nivel", "cantidad"]
DISP_COLS = ["grupo", "fecha", "estado", "concepto", "moneda", "tipo_moneda", "importe"]
BANCOS_COLS = ["grupo", "fecha", "seccion", "empresa", "banco", "cuenta", "moneda", "importe_ars", "importe_moneda"]
TC_COLS = ["fecha", "ars_por_usd"]

TABLAS = {
    "hf": (HF_COLS, ("hotel", "fecha")),
    "flash": (FLASH_COLS, ("hotel", "fecha", "concepto")),
    "pickup": (PICKUP_COLS, ("hotel", "fecha_reporte", "mes")),
    "bonvoy": (BONVOY_COLS, ("hotel", "fecha", "nivel")),
    "tipo_cambio": (TC_COLS, ("fecha",)),
}


def ruta(tabla: str) -> Path:
    return DATA / f"{tabla}.csv"


def leer(tabla: str) -> list[dict]:
    r = ruta(tabla)
    if not r.exists() or r.stat().st_size == 0:
        return []
    with r.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def escribir(tabla: str, columnas: list[str], filas: list[dict]) -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    with ruta(tabla).open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columnas, extrasaction="ignore")
        w.writeheader()
        w.writerows(filas)


def fmt(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, (date, datetime)):
        return v.isoformat()[:10]
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.6f}".rstrip("0").rstrip(".") if abs(v) < 10 else f"{v:.2f}"
    return str(v)


def _norm(fila: dict, columnas: list[str]) -> dict:
    return {c: fmt(fila.get(c)) for c in columnas}


def upsert(tabla: str, filas: list[dict], reemplazar_por: tuple[str, ...] | None = None) -> int:
    """Inserta o reemplaza por clave. Si una fila existente tiene fecha_reporte posterior, se conserva.

    reemplazar_por: borra antes todas las filas existentes que coincidan en esos campos con alguna nueva
    (para tablas donde un reporte trae el conjunto completo, p. ej. bonvoy de un día).
    """
    columnas, clave = TABLAS[tabla]
    nuevas = [_norm(f, columnas) for f in filas]
    existentes = leer(tabla)
    if reemplazar_por:
        grupos = {tuple(f[c] for c in reemplazar_por) for f in nuevas}
        existentes = [f for f in existentes if tuple(f[c] for c in reemplazar_por) not in grupos]
    indice = {tuple(f[c] for c in clave): f for f in existentes}
    for f in nuevas:
        k = tuple(f[c] for c in clave)
        viejo = indice.get(k)
        if viejo and "fecha_reporte" in columnas and viejo.get("fecha_reporte", "") > f.get("fecha_reporte", ""):
            continue
        indice[k] = f
    escribir(tabla, columnas, sorted(indice.values(), key=lambda f: tuple(f[c] for c in clave)))
    return len(nuevas)


def reemplazar_bloque(tabla: str, columnas: list[str], filas: list[dict], clave: tuple[str, ...]) -> int:
    """Para disponibilidades y bancos: el archivo de un grupo y fecha reemplaza todo lo anterior de ese grupo y fecha."""
    nuevas = [_norm(f, columnas) for f in filas]
    bloques = {tuple(f[c] for c in clave) for f in nuevas}
    resto = [f for f in leer(tabla) if tuple(f[c] for c in clave) not in bloques]
    escribir(tabla, columnas, sorted(resto + nuevas, key=lambda f: tuple(f[c] for c in clave)))
    return len(nuevas)


# ---------------------------------------------------------------- registro de mails procesados


def _registro() -> dict:
    if REGISTRO.exists():
        return json.loads(REGISTRO.read_text(encoding="utf-8"))
    return {}


def ya_procesado(clave: str) -> bool:
    return clave in _registro()


def marcar_procesado(claves: dict[str, dict]) -> None:
    reg = _registro()
    reg.update(claves)
    DATA.mkdir(parents=True, exist_ok=True)
    REGISTRO.write_text(json.dumps(reg, ensure_ascii=False, indent=1, sort_keys=True), encoding="utf-8")


def ahora() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="minutes")


# ---------------------------------------------------------------- Excel para Power BI


def exportar_excel(hoteles: list[dict]) -> None:
    """data/hoteles.xlsx con las mismas columnas que las hojas de la base histórica."""
    from openpyxl import Workbook

    base = {h["id"]: h["base"] for h in hoteles}
    dias = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]

    def d(s: str):
        return datetime.strptime(s, "%Y-%m-%d").date() if s else None

    def n(s: str):
        return float(s) if s not in ("", None) else None

    wb = Workbook()
    ws = wb.active
    ws.title = "H&F"
    ws.append(["Date", "Total\nOcc.", "Arr.\nRooms", "Comp.\nRooms", "House\nUse", "Deduct\nIndiv.", "Deduct\nGroup",
               "Occ.%", "Room Revenue", "Average Rate", "Dep.\nRooms", "Day Use\nRooms", "No Show\nRooms",
               "OOO\nRooms", "Adl. &\nChl.", "Fecha", "Día", "HoF", "Empresa"])
    for f in leer("hf"):
        fe = d(f["fecha"])
        ws.append([fe.strftime("%d-%m-%y %a"), *[n(f[c]) for c in HF_COLS[3:17]], fe, dias[fe.weekday()],
                   f["tipo"], base.get(f["hotel"], f["hotel"])])

    ws = wb.create_sheet("Flash")
    ws.append(["CONCEPTO", "DAY", "MONTH", "YEAR", "MONEDA", "DATE", "HOTEL", "FECHA NEGOCIO"])
    for f in leer("flash"):
        ws.append([f["concepto"], n(f["dia"]), n(f["mes"]), n(f["anio"]), "USD", d(f["fecha_reporte"]),
                   base.get(f["hotel"], f["hotel"]), d(f["fecha"])])

    ws = wb.create_sheet("Pick up")
    ws.append(["Fecha", "Mes", "% Occ", "NTS", "GRP", "ADR", "Revenue", "empresa"])
    for f in leer("pickup"):
        ws.append([d(f["fecha_reporte"]), f["mes"], n(f["occ_pct"]), n(f["noches"]), n(f["grupo"]), n(f["adr"]),
                   n(f["revenue"]), base.get(f["hotel"], f["hotel"])])

    ws = wb.create_sheet("Bonvoy")
    ws.append(["Bonvoy", "Cantidad", "Fecha", "Empresa"])
    for f in leer("bonvoy"):
        ws.append([f["nivel"], n(f["cantidad"]), d(f["fecha"]), base.get(f["hotel"], f["hotel"])])

    ws = wb.create_sheet("Disponibilidades")
    ws.append(["Date", "Estado", "CONCEPTO", "Moneda", "Tipo de moneda", "Importe", "Grupo"])
    for f in leer("disponibilidades"):
        ws.append([d(f["fecha"]), f["estado"], f["concepto"], f["moneda"], f["tipo_moneda"], n(f["importe"]),
                   f["grupo"]])

    ws = wb.create_sheet("Países")
    ws.append(["Mes", "PAÍS", "Continente", "Huéspedes", "Empresa"])
    for f in leer("paises"):
        ws.append([f["mes"], f["pais"], f["continente"], n(f["huespedes"]), base.get(f["hotel"], f["hotel"])])

    ws = wb.create_sheet("Tipo de cambio")
    ws.append(["Fecha", "BNA vendedor"])
    for f in leer("tipo_cambio"):
        ws.append([d(f["fecha"]), n(f["ars_por_usd"])])
    wb.save(EXCEL)
