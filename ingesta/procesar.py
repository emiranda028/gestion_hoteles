"""Clasifica cada adjunto (PDF, ZIP o Excel), lo lee y guarda el resultado en data/."""
from __future__ import annotations

import calendar
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, timedelta

from . import almacen, disponibilidades, opera

# Conceptos del Manager Flash que se guardan (el resto del reporte no se usa en tableros)
CONCEPTOS_FLASH = [
    "Total Rooms in Hotel", "Total Rooms in Hotel minus OOO Rooms", "Rooms Occupied",
    "Rooms Occupied minus House Use", "Complimentary Rooms", "House Use Rooms", "Out of Order Rooms",
    "Day Use Rooms", "Total In-House Persons", "Individual Rooms In-House", "Block Rooms In-House",
    "Member Persons In-House", "Arrival Rooms", "Departure Rooms", "No Show Rooms",
    "Cancelled Reservations for Today", "Reservations Made Today", "Reservation Cancellations made Today",
    "Room Nights Reserved Today", "% Rooms Occupied", "ADR", "Room Revenue", "Food And Beverage Revenue",
    "Other Revenue", "Total Revenue", "Member Room Revenue", "Revenue per Available Room",
    "% Rooms Occupied for Tomorrow", "% Rooms Occupied for the Next 7 Days", "% Rooms Occupied for the Next 31 Days",
]

NIVELES_BONVOY = [
    ("AMBASSADOR", "Ambassador Elite (AMB)"), ("TITANIUM", "Titanium Elite (TTM)"),
    ("PLATINUM", "Platinum Elite (PLT)"), ("GOLD", "Gold Elite (GLD)"), ("SILVER", "Silver Elite (SLR)"),
    ("MEMBER", "Member (MRD)"),
]


def nivel_bonvoy(texto: str) -> str:
    t = texto.upper()
    for clave, nombre in NIVELES_BONVOY:
        if clave in t:
            return nombre
    return texto.strip()


@dataclass
class Resultado:
    origen: str
    tipo: str
    ok: bool
    hotel: str | None = None
    fecha: str | None = None
    detalle: str = ""
    avisos: list[str] = field(default_factory=list)


def hotel_por_nombre(nombre: str, config: dict) -> dict | None:
    n = nombre.lower()
    for h in config["hoteles"]:
        if h["opera"].lower() in n:
            return h
    return None


def grupo_por_empresa(empresa: str, config: dict) -> dict | None:
    e = empresa.upper()
    for g in config["grupos"]:
        if any(x.upper() in e for x in g["detectar"]):
            return g
    return None


def expandir(nombre: str, contenido: bytes) -> list[tuple[str, bytes]]:
    """Los reportes suelen venir dentro de un .zip: se devuelven los archivos internos."""
    if nombre.lower().endswith(".zip"):
        salida = []
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            for info in z.infolist():
                if not info.is_dir():
                    salida += expandir(info.filename.split("/")[-1], z.read(info))
        return salida
    return [(nombre, contenido)]


def procesar(nombre: str, contenido: bytes, config: dict, fecha_recepcion: date | None) -> list[Resultado]:
    resultados = []
    for n, c in expandir(nombre, contenido):
        low = n.lower()
        try:
            if low.endswith(".pdf"):
                resultados.append(_pdf(n, c, config))
            elif low.endswith((".xlsx", ".xlsm")):
                r = _excel(n, c, config, fecha_recepcion)
                if r:
                    resultados.append(r)
        except Exception as e:  # un archivo dañado no frena al resto
            resultados.append(Resultado(n, "?", False, avisos=[f"No se pudo leer: {e}"]))
    return resultados


def _pdf(nombre: str, contenido: bytes, config: dict) -> Resultado:
    texto = opera.extraer_texto(contenido)
    tipo = opera.tipo_reporte(texto)
    if tipo is None:
        return Resultado(nombre, "ignorado", True, detalle="no es un reporte que se use en los tableros")
    if tipo == "auditoria":
        return _auditoria(nombre, texto, config)
    nombre_hotel, _ = opera.cabecera(texto)
    hotel = hotel_por_nombre(nombre_hotel, config)
    if hotel is None:
        return Resultado(nombre, tipo, False, avisos=[f"Hotel no configurado: '{nombre_hotel}' (agregarlo en config.yaml)"])
    hid = hotel["id"]

    if tipo == "flash":
        f = opera.leer_flash(texto)
        if f.moneda != "USD":
            return Resultado(nombre, tipo, True, hid, f.fecha_negocio.isoformat(),
                             avisos=[f"Flash en {f.moneda}: se ignora, se usa solo el de USD"])
        filas = []
        for concepto in CONCEPTOS_FLASH:
            v = f.conceptos.get(concepto)
            if v:
                filas.append(_fila_flash(hid, f, concepto, v))
        filas += _derivados_flash(hid, f, hotel)
        faltan = [c for c in ("Rooms Occupied", "Room Revenue", "Total Revenue") if c not in f.conceptos]
        almacen.upsert("flash", filas)
        return Resultado(nombre, "Manager Flash", not faltan, hid, f.fecha_negocio.isoformat(),
                         detalle=f"{len(filas)} conceptos",
                         avisos=[f"Faltan conceptos: {', '.join(faltan)}"] if faltan else [])

    if tipo == "hf":
        r = opera.leer_hf(texto)
        if r.moneda != "USD":
            return Resultado(nombre, "History & Forecast", True, hid, r.hasta.isoformat(),
                             avisos=[f"H&F en {r.moneda}: se ignora (se usan los reportes en USD)"])
        emitido = (r.fecha_reporte or date.today()).isoformat()
        filas = [{**d.__dict__, "hotel": hid, "fecha_reporte": emitido} for d in r.dias]
        almacen.upsert("hf", filas)
        if r.mes_completo and r.total:
            t = r.total
            almacen.upsert("pickup", [{
                "hotel": hid, "fecha_reporte": emitido, "mes": r.desde.isoformat()[:7], "noches": t.total_occ,
                "grupo": t.deduct_group, "occ_pct": t.occ_pct, "revenue": t.room_revenue, "adr": t.adr,
            }])
        return Resultado(nombre, "History & Forecast", True, hid, f"{r.desde} a {r.hasta}",
                         detalle=f"{len(filas)} días")

    e = opera.leer_elite(texto)
    filas = [{"hotel": hid, "fecha": f, "nivel": nivel_bonvoy(n), "cantidad": c} for (f, n), c in e.llegadas.items()]
    if filas:
        almacen.upsert("bonvoy", filas, reemplazar_por=("hotel", "fecha"))
    fechas = sorted({f.isoformat() for f, _ in e.llegadas})
    return Resultado(nombre, "Elite Arrivals", True, hid, ", ".join(fechas),
                     detalle=f"{sum(e.llegadas.values())} llegadas Bonvoy")


def _auditoria(nombre: str, texto: str, config: dict) -> Resultado:
    """La auditoría no trae el nombre del hotel: la manda solo el hotel indicado en config (auditoria_hotel)."""
    hid = config.get("auditoria_hotel")
    if not hid:
        return Resultado(nombre, "Auditoría", False, avisos=["Falta 'auditoria_hotel' en config.yaml"])
    a = opera.leer_auditoria(texto)
    # igual que en la base: la fecha es la del día siguiente a la auditoría (día de llegada)
    fecha = a.fecha + timedelta(days=1)
    filas = [{"hotel": hid, "fecha": fecha, "nivel": nivel_bonvoy(n), "cantidad": c} for n, c in a.niveles.items()]
    if filas:
        almacen.upsert("bonvoy", filas, reemplazar_por=("hotel", "fecha"))
    return Resultado(nombre, "Auditoría (Membership)", True, hid, fecha.isoformat(),
                     detalle=f"{sum(a.niveles.values())} huéspedes Bonvoy")


def _fila_flash(hid: str, f: opera.Flash, concepto: str, v: tuple) -> dict:
    return {"hotel": hid, "fecha": f.fecha_negocio, "fecha_reporte": f.fecha_reporte, "concepto": concepto,
            "dia": v[0], "mes": v[1], "anio": v[2]}


def _derivados_flash(hid: str, f: opera.Flash, hotel: dict) -> list[dict]:
    """Conceptos calculados que usa la base de Power BI (hoja Flash)."""
    c = f.conceptos
    cero = (0.0, 0.0, 0.0)
    total = c.get("Total Rooms in Hotel", cero)
    ocup = c.get("Rooms Occupied minus House Use", cero)
    personas = c.get("Total In-House Persons", cero)
    rev = c.get("Room Revenue", cero)

    def div(a, b):
        return tuple((x / y) if x is not None and y else None for x, y in zip(a, b))

    h = float(hotel.get("habitaciones") or 0)
    salida = {
        "Habitaciones Construidas": (h, h, h),
        "Tasa Ocupación": div(ocup, total),
        "Tasa Doble Ocupación": div(personas, ocup),
        "Ventas Totales": c.get("Total Revenue", cero),
        "RevPar": div(rev, total),
    }
    return [_fila_flash(hid, f, k, v) for k, v in salida.items()]


def _excel(nombre: str, contenido: bytes, config: dict, fecha: date | None) -> Resultado | None:
    from openpyxl import load_workbook

    from . import paises

    hojas = load_workbook(io.BytesIO(contenido), read_only=True).sheetnames
    if not disponibilidades.es_planilla(hojas):
        if paises.es_planilla(contenido):
            filas = paises.leer(contenido)
            paises.guardar(filas)
            return Resultado(nombre, "Huéspedes por país", True, paises.HOTEL, max(f["mes"] for f in filas),
                             detalle=f"{len(filas)} filas")
        return None  # otros Excel (conversiones manuales, etc.) no se usan
    d = disponibilidades.leer(contenido)
    grupo = grupo_por_empresa(d.empresa, config)
    if grupo is None:
        return Resultado(nombre, "Disponibilidades", False, avisos=[f"Empresa no configurada: {d.empresa}"])
    # Convención de la base: la fecha es la del día en que llega el informe
    fecha = fecha or d.fecha_saldos or date.today()
    gid = grupo["id"]
    almacen.reemplazar_bloque("disponibilidades", almacen.DISP_COLS,
                              [{"grupo": gid, "fecha": fecha, **x} for x in d.resumen], ("grupo", "fecha"))
    almacen.reemplazar_bloque("bancos", almacen.BANCOS_COLS,
                              [{"grupo": gid, "fecha": fecha, **x} for x in d.detalle], ("grupo", "fecha"))
    avisos = []
    if d.tc_bna and not any(r["fecha"] == fecha.isoformat() for r in almacen.leer("tipo_cambio")):
        almacen.upsert("tipo_cambio", [{"fecha": fecha, "ars_por_usd": d.tc_bna}])
    total = next(x["importe"] for x in d.resumen if x["concepto"] == "DISPONIBILIDADES")
    if not total:
        avisos.append("El total de disponibilidades es cero")
    return Resultado(nombre, "Disponibilidades", not avisos, gid, fecha.isoformat(),
                     detalle=f"total $ {total:,.0f}".replace(",", "."), avisos=avisos)


def ultimo_dia_mes(d: date) -> int:
    return calendar.monthrange(d.year, d.month)[1]
