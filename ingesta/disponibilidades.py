"""Lectura del Excel diario de disponibilidades (hoja "Disponibilidades Anexo") que mandan los grupos.

Devuelve el resumen con los mismos conceptos que la hoja "Disponibilidades PBI" de la base
histórica y el detalle por cuenta. Los importes quedan como vienen (pesos, o la moneda original
para dólares y euros); la conversión a dólares la hace la web app con el BNA vendedor del día.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

HOJA = "Disponibilidades Anexo"

# columnas de la hoja (0 = A)
C_SECCION, C_FECHA, C_EMPRESA, C_BANCO, C_CUENTA, C_IMPORTE, C_EXTRACTO, C_TC = 1, 2, 3, 4, 5, 6, 7, 8


@dataclass
class Disponibilidades:
    empresa: str
    fecha_saldos: date | None
    resumen: list[dict] = field(default_factory=list)  # estado, concepto, moneda, tipo_moneda, importe
    detalle: list[dict] = field(default_factory=list)  # seccion, empresa, banco, cuenta, moneda, importe_ars, importe_moneda
    tc_bna: float | None = None  # si la planilla trae "TP BNA Billete venta"


def es_planilla(nombre_hojas: list[str]) -> bool:
    return HOJA in nombre_hojas


def _txt(v) -> str:
    return str(v).strip() if v not in (None, 0) else ""


def _num(v) -> float:
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def leer(contenido: bytes) -> Disponibilidades:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    filas = [list(r) + [None] * 12 for r in wb[HOJA].iter_rows(values_only=True)]
    return leer_filas(filas)


def leer_filas(filas: list[list]) -> Disponibilidades:
    seccion = ""
    empresa = ""
    fecha_saldos = None
    tc_bna = None
    t = {k: 0.0 for k in ("bancos", "me_ars", "inversiones", "cheques", "pagos", "recaudacion", "cobros",
                          "socios", "total", "tc_usd", "tc_eur")}
    me = {"USD": 0.0, "EUR": 0.0}
    efectivo: dict[tuple[str, str], list[float]] = {}  # (Tesorería|Seguridad, USD|EUR) -> [ars, moneda]
    bancos_usd = [0.0, 0.0]
    detalle: list[dict] = []
    total_visto = False
    rubro = ""

    for r in filas:
        s = _txt(r[C_SECCION]).upper()
        etiqueta = _txt(r[C_EMPRESA])
        et = etiqueta.upper()
        if s in ("BANCOS", "MONEDA EXTRANJERA", "INVERSIONES"):
            seccion = s
            continue
        texto_fila = " ".join(_txt(x) for x in r[:6]).upper()
        if "PAGOS PROYECTADOS" in texto_fila and s != "SUB-TOTAL":
            seccion = "PAGOS"
        elif "COBROS PROYECTADO" in texto_fila and s != "SUB-TOTAL":
            seccion = "COBROS"
        elif s == "ACCIONISTAS" or "SOCIOS" == s:
            seccion = "SOCIOS"
        if s and s not in ("SUB-TOTAL", "TOTAL"):
            rubro = s.title()
        if isinstance(r[C_FECHA], datetime) and fecha_saldos is None and seccion == "BANCOS":
            fecha_saldos = r[C_FECHA].date()

        if s == "TOTAL" and not total_visto:
            t["total"] = _num(r[C_IMPORTE])
            total_visto = True
            continue
        if s == "SUB-TOTAL":
            if seccion == "PAGOS" and "PAGOS" in et:
                t["pagos"] = _num(r[C_IMPORTE])
            elif seccion == "COBROS" and "COBROS" in et:
                t["cobros"] = _num(r[C_IMPORTE])
            elif seccion == "SOCIOS":
                t["socios"] = _num(r[C_IMPORTE])
            continue

        if seccion == "BANCOS":
            if et.startswith("TOTAL BANCOS"):
                t["bancos"] = _num(r[C_IMPORTE])
            elif etiqueta and _txt(r[C_BANCO]):
                empresa = empresa or etiqueta
                detalle.append({"seccion": "Bancos pesos", "empresa": etiqueta, "banco": _txt(r[C_BANCO]),
                                "cuenta": _txt(r[C_CUENTA]), "moneda": "ARS", "importe_ars": _num(r[C_IMPORTE]),
                                "importe_moneda": _num(r[C_IMPORTE])})
        elif seccion == "MONEDA EXTRANJERA":
            banco = _txt(r[C_BANCO])
            if et.startswith("TOTAL USD"):
                t["tc_usd"] = _num(r[C_TC])
            elif et.startswith("TOTAL EUR"):
                t["tc_eur"] = _num(r[C_TC])
            elif et.startswith("TOTAL MONEDA"):
                t["me_ars"] = _num(r[C_IMPORTE])
            elif etiqueta and banco:
                moneda = "EUR" if "EUR" in banco.upper() else "USD"
                ars, orig = _num(r[C_IMPORTE]), _num(r[C_EXTRACTO])
                me[moneda] += orig
                if "CAJA" in banco.upper() or "EFECTIVO" in banco.upper():
                    lugar = "Seguridad" if "SEGURIDAD" in _txt(r[C_CUENTA]).upper() else "Tesorería"
                    x = efectivo.setdefault((lugar, moneda), [0.0, 0.0])
                    x[0] += ars
                    x[1] += orig
                elif moneda == "USD":
                    bancos_usd[0] += ars
                    bancos_usd[1] += orig
                if moneda == "USD" and _num(r[C_TC]) and "BNA" in str(
                        filas[_idx_cabecera(filas, "MONEDA EXTRANJERA")][C_TC] or "").upper():
                    tc_bna = _num(r[C_TC])
                detalle.append({"seccion": "Moneda extranjera", "empresa": etiqueta, "banco": banco,
                                "cuenta": _txt(r[C_CUENTA]), "moneda": moneda, "importe_ars": ars,
                                "importe_moneda": orig})
        elif seccion == "INVERSIONES":
            if et.startswith("TOTAL FONDO"):
                t["inversiones"] = _num(r[C_IMPORTE])
            elif etiqueta and _txt(r[C_BANCO]):
                detalle.append({"seccion": "Inversiones", "empresa": etiqueta, "banco": _txt(r[C_BANCO]),
                                "cuenta": _txt(r[C_CUENTA]), "moneda": "ARS", "importe_ars": _num(r[C_IMPORTE]),
                                "importe_moneda": _num(r[C_IMPORTE])})
        elif seccion == "PAGOS":
            if "CHEQUES EMITIDOS" in et:
                t["cheques"] += _num(r[C_IMPORTE])
            if etiqueta and _num(r[C_IMPORTE]):
                detalle.append({"seccion": "Pagos proyectados", "empresa": _txt(r[C_BANCO]) or _txt(r[C_CUENTA]),
                                "banco": rubro, "cuenta": etiqueta, "moneda": "ARS",
                                "importe_ars": -_num(r[C_IMPORTE]), "importe_moneda": -_num(r[C_IMPORTE])})
        elif seccion == "COBROS":
            if "RECAUDACI" in et:
                t["recaudacion"] += _num(r[C_IMPORTE])
            valor = _num(r[C_IMPORTE]) or _num(r[C_BANCO]) or _num(r[C_CUENTA])
            if etiqueta and valor:
                detalle.append({"seccion": "Cobros proyectados", "empresa": _txt(r[C_CUENTA]) or _txt(r[C_BANCO]),
                                "banco": rubro, "cuenta": etiqueta, "moneda": "ARS",
                                "importe_ars": valor, "importe_moneda": valor})

    if not empresa:
        raise ValueError("No se encontró la empresa en la sección BANCOS")

    R = "Real"
    P = "Proyectado"
    resumen = [
        (R, "DISPONIBILIDADES", "Local", "ARS", t["total"]),
        (R, "Moneda Local", "Local", "ARS", t["total"] - t["me_ars"]),
        (R, "Moneda extranjera", "Local", "ARS", t["me_ars"]),
        (R, "Moneda extranjera", "Extranjera", "USD", me["USD"]),
        (R, "Tipo de cambio USD", "Local", "USD", t["tc_usd"]),
        (R, "Moneda extranjera", "Extranjera", "EUR", me["EUR"]),
        (R, "Tipo de cambio EUR", "Local", "EUR", t["tc_eur"]),
        (R, "Bancos dólares", "Extranjera", "USD", bancos_usd[1]),
        (R, "Bancos dólares", "Local", "ARS", bancos_usd[0]),
        (R, "Bancos pesos", "Local", "ARS", t["bancos"]),
        (R, "Inversiones", "Local", "ARS", t["inversiones"]),
    ]
    for (lugar, moneda), (ars, orig) in sorted(efectivo.items()):
        resumen.append((R, f"Efectivo - {lugar}", "Extranjera", moneda, orig))
        resumen.append((R, f"Efectivo - {lugar}", "Local", "ARS", ars))
    resumen += [
        (P, "Efectivo - Recaudación", "Local", "ARS", t["recaudacion"]),
        (P, "Cobranzas Proyectadas", "Local", "ARS", t["cobros"] - t["recaudacion"]),
        (P, "Cheques emitidos", "Local", "ARS", -t["cheques"]),
        (P, "Pagos programados", "Local", "ARS", -(t["pagos"] - t["cheques"])),
        (P, "Aportes socios", "Local", "ARS", t["socios"]),
    ]
    return Disponibilidades(
        empresa=empresa,
        fecha_saldos=fecha_saldos,
        resumen=[dict(zip(("estado", "concepto", "moneda", "tipo_moneda", "importe"), x)) for x in resumen],
        detalle=detalle,
        tc_bna=tc_bna,
    )


def _idx_cabecera(filas: list[list], nombre: str) -> int:
    for i, r in enumerate(filas):
        if _txt(r[C_SECCION]).upper() == nombre:
            return i
    return 0
