"""Punto de entrada de la ingesta.

Uso:
  python -m ingesta.run gmail                      # baja PDFs nuevos de Gmail y actualiza data/
  python -m ingesta.run carpeta RUTA               # procesa PDFs locales (carga histórica)
  python -m ingesta.run inspeccionar ARCHIVO.pdf   # muestra texto y campos detectados (para ajustar config)
  python -m ingesta.run excel ARCHIVO.xlsx [--hoja NOMBRE]   # importa el Excel histórico
"""
from __future__ import annotations

import argparse
import json
import sys
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path

import yaml

from . import almacen
from .parser import CAMPOS, Reporte, extraer_texto, leer_reporte

CONFIG = Path(__file__).with_name("config.yaml")
ESTADO = almacen.DATA / "ultima_ingesta.json"


def cargar_config() -> dict:
    return yaml.safe_load(CONFIG.read_text(encoding="utf-8"))


def _resumen(rep: Reporte, origen: str) -> dict:
    return {
        "origen": origen,
        "hotel": rep.hotel_id,
        "fecha": rep.fecha.isoformat() if rep.fecha else None,
        "ok": rep.valido,
        "avisos": rep.avisos,
    }


def _guardar_estado(resultados: list[dict]) -> None:
    ESTADO.parent.mkdir(parents=True, exist_ok=True)
    ESTADO.write_text(
        json.dumps(
            {"ejecutado": datetime.now(timezone.utc).isoformat(timespec="minutes"), "archivos": resultados},
            ensure_ascii=False,
            indent=1,
        ),
        encoding="utf-8",
    )


def _imprimir(resultados: list[dict]) -> int:
    fallidos = [r for r in resultados if not r["ok"]]
    for r in resultados:
        marca = "OK " if r["ok"] else "ERR"
        print(f"[{marca}] {r['fecha']} {r['hotel']} <- {r['origen']}")
        for a in r["avisos"]:
            print(f"       ! {a}")
    print(f"\n{len(resultados) - len(fallidos)} procesados, {len(fallidos)} con error.")
    return 1 if fallidos else 0


def cmd_gmail(_args) -> int:
    from .gmail import Gmail

    config = cargar_config()
    cliente = Gmail(config["gmail"])
    reportes, origenes, resultados, validos, uids_ok = [], [], [], [], set()
    try:
        for adj in cliente.adjuntos_pdf():
            clave = f"{adj.message_id}|{almacen.huella(adj.contenido)}"
            if almacen.ya_procesado(clave):
                continue
            contexto = f"{adj.asunto}\n{adj.remitente}\n{adj.archivo}"
            origen = f"gmail:{adj.fecha_mail} {adj.archivo}"
            try:
                rep = leer_reporte(adj.contenido, config, contexto, adj.fecha_mail)
            except Exception as e:  # PDF dañado o ilegible: se informa y se sigue
                resultados.append({"origen": origen, "hotel": None, "fecha": None, "ok": False,
                                   "avisos": [f"No se pudo leer el PDF: {e}"]})
                continue
            reportes.append(rep)
            origenes.append(origen)
            resultados.append(_resumen(rep, origen))
            if rep.valido:
                validos.append((clave, {"archivo": adj.archivo, "hotel": rep.hotel_id,
                                        "fecha": rep.fecha.isoformat()}))
                uids_ok.add(adj.uid)
        almacen.guardar(reportes, config["hoteles"], origenes)
        for clave, info in validos:
            almacen.marcar_procesado(clave, info)
        for uid in uids_ok:
            cliente.etiquetar(uid)
    finally:
        cliente.cerrar()
    _guardar_estado(resultados)
    if not resultados:
        print("No hay reportes nuevos.")
        return 0
    return _imprimir(resultados)


def cmd_carpeta(args) -> int:
    config = cargar_config()
    archivos = sorted(Path(args.ruta).rglob("*.pdf")) + sorted(Path(args.ruta).rglob("*.PDF"))
    reportes, origenes, resultados = [], [], []
    for ruta in archivos:
        fecha_archivo = date.fromtimestamp(ruta.stat().st_mtime)
        try:
            rep = leer_reporte(ruta.read_bytes(), config, ruta.name, fecha_archivo)
        except Exception as e:
            resultados.append({"origen": ruta.name, "hotel": None, "fecha": None, "ok": False,
                               "avisos": [f"No se pudo leer el PDF: {e}"]})
            continue
        reportes.append(rep)
        origenes.append(f"archivo:{ruta.name}")
        resultados.append(_resumen(rep, ruta.name))
    almacen.guardar(reportes, config["hoteles"], origenes)
    _guardar_estado(resultados)
    return _imprimir(resultados)


def cmd_inspeccionar(args) -> int:
    config = cargar_config()
    contenido = Path(args.archivo).read_bytes()
    texto = extraer_texto(contenido)
    print("=" * 30, "TEXTO EXTRAÍDO", "=" * 30)
    print(texto)
    rep = leer_reporte(contenido, config, Path(args.archivo).name, date.today(), texto=texto)
    print("=" * 30, "RESULTADO", "=" * 30)
    print(f"plantilla: {rep.plantilla}\nhotel:     {rep.hotel_id}\nfecha:     {rep.fecha}")
    for c in CAMPOS:
        v = rep.valores.get(c)
        print(f"  {c:28s} {'—' if v is None else v}")
    if rep.procedencia:
        print("procedencia:")
        for p in rep.procedencia:
            print(f"  {p['pais']:24s} {p['iso2'] or '??'} {p['huespedes']}")
    for a in rep.avisos:
        print(f"! {a}")
    print("\nVÁLIDO" if rep.valido else "\nINCOMPLETO: ajustar patrones en ingesta/config.yaml")
    return 0 if rep.valido else 1


# --- Importación del Excel histórico -------------------------------------------------

SINONIMOS = {
    "fecha": ["fecha", "dia", "date", "fecha de cierre"],
    "hotel": ["hotel", "establecimiento", "propiedad"],
    "habitaciones_disponibles": ["habitaciones disponibles", "hab disponibles", "disponibles", "rooms available"],
    "habitaciones_ocupadas": ["habitaciones ocupadas", "hab ocupadas", "ocupadas", "rooms occupied", "rooms sold"],
    "habitaciones_fuera_servicio": ["fuera de servicio", "ooo", "out of order"],
    "huespedes": ["huespedes", "pax", "pasajeros", "guests"],
    "llegadas": ["llegadas", "arrivals", "check in"],
    "salidas": ["salidas", "departures", "check out"],
    "ocupacion_pct": ["ocupacion", "% ocupacion", "ocupacion %", "occupancy"],
    "adr": ["adr", "tarifa promedio", "tarifa media"],
    "ingreso_habitaciones": ["ingreso habitaciones", "ingresos habitaciones", "ingreso alojamiento", "room revenue"],
    "ingreso_ayb": ["ingreso ayb", "ingresos ayb", "ingreso a&b", "alimentos y bebidas", "f&b revenue"],
    "ingreso_otros": ["otros ingresos", "ingreso otros", "other revenue"],
    "ingreso_total": ["ingreso total", "ingresos totales", "total ingresos", "total revenue"],
}


def _norm(t) -> str:
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return " ".join(t.replace("_", " ").replace(".", " ").split())


def cmd_excel(args) -> int:
    from openpyxl import load_workbook

    from .numeros import a_numero
    from .parser import _completar_derivados, _parsear_fecha, detectar_hotel

    config = cargar_config()
    wb = load_workbook(args.archivo, data_only=True, read_only=True)
    ws = wb[args.hoja] if args.hoja else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    encabezado = [_norm(c) for c in next(filas)]
    mapa = {}
    for campo, sinonimos in SINONIMOS.items():
        for i, col in enumerate(encabezado):
            if col in sinonimos:
                mapa[campo] = i
                break
    print("Columnas reconocidas:", {k: encabezado[v] for k, v in mapa.items()})
    if "fecha" not in mapa:
        print("No se encontró la columna de fecha.")
        return 1

    reportes, origenes = [], []
    for fila in filas:
        crudo = fila[mapa["fecha"]]
        f = crudo.date() if isinstance(crudo, datetime) else crudo if isinstance(crudo, date) else (
            _parsear_fecha(str(crudo)) if crudo else None)
        if f is None:
            continue
        hotel = (detectar_hotel(str(fila[mapa["hotel"]]), config["hoteles"]) if "hotel" in mapa else None)
        if hotel is None and args.hotel:
            hotel = next((h for h in config["hoteles"] if h["id"] == args.hotel), None)
        valores = {}
        for c in CAMPOS:
            v = fila[mapa[c]] if c in mapa else None
            valores[c] = float(v) if isinstance(v, (int, float)) else a_numero(v)
        if valores["ocupacion_pct"] is not None and valores["ocupacion_pct"] <= 1:
            valores["ocupacion_pct"] *= 100  # Excel guarda 85% como 0,85
        avisos: list[str] = []
        _completar_derivados(valores, hotel, avisos)
        reportes.append(Reporte(hotel["id"] if hotel else None, f, "excel", valores, [], avisos))
        origenes.append(f"excel:{Path(args.archivo).name}")
    n = almacen.guardar(reportes, config["hoteles"], origenes)
    sin_hotel = sum(1 for r in reportes if r.hotel_id is None)
    print(f"{n} días importados. {sin_hotel} filas sin hotel identificado"
          + (" (usar --hotel ID si el Excel es de un solo hotel)." if sin_hotel else "."))
    return 0


def main(argv=None) -> int:
    p = argparse.ArgumentParser(prog="python -m ingesta.run", description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("gmail").set_defaults(fn=cmd_gmail)
    c = sub.add_parser("carpeta")
    c.add_argument("ruta")
    c.set_defaults(fn=cmd_carpeta)
    i = sub.add_parser("inspeccionar")
    i.add_argument("archivo")
    i.set_defaults(fn=cmd_inspeccionar)
    e = sub.add_parser("excel")
    e.add_argument("archivo")
    e.add_argument("--hoja")
    e.add_argument("--hotel", help="id del hotel si el Excel no tiene columna Hotel")
    e.set_defaults(fn=cmd_excel)
    args = p.parse_args(argv)
    return args.fn(args)


if __name__ == "__main__":
    sys.exit(main())
